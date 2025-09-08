from typing import Dict
from datetime import datetime
import openpyxl
from app.utils.db import get_connection
from app.utils.id_utils import generate_game_data_id

# gameData 表（按 db.mdc 扩展需求）：
# id, gameID, date, category, channel, daNewUser, newUser, active, payUser, gain, ARPU, ARPPU, DAY1


def ensure_game_data_table():
    sql = (
        'CREATE TABLE IF NOT EXISTS `gameData` ('
        '  `id` VARCHAR(32) PRIMARY KEY,'
        '  `gameID` VARCHAR(32) NOT NULL,'
        '  `date` DATETIME,'
        '  `category` VARCHAR(64),'
        '  `channel` VARCHAR(128),'
        '  `daNewUser` INT,'
        '  `newUser` INT,'
        '  `active` INT,'
        '  `payUser` INT,'
        '  `gain` DOUBLE,'
        '  `ARPU` DOUBLE,'
        '  `ARPPU` DOUBLE,'
        '  `DAY1` DOUBLE,'
        '  `DAY7` DOUBLE'
        ') ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;'
    )
    conn = get_connection()
    with conn:
        with conn.cursor() as cur:
            cur.execute(sql)
        conn.commit()


HEADER_TO_FIELD = {
    '日期': 'date',
    '渠道': 'channel',
    '新增用户': 'newUser',
    'DA新增用户': 'daNewUser',
    '活跃用户': 'active',
    '新增留存_新增次留率': 'DAY1',
    '新增留存_新增7留率': 'DAY7',
    '付费数据_付费用户': 'payUser',
    '付费数据_收入': 'gain',
    '付费数据_ARPU': 'ARPU',
    '付费数据_ARPPU': 'ARPPU',
}


def parse_date(value) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d'):
            try:
                return datetime.strptime(value.strip(), fmt)
            except ValueError:
                continue
    raise ValueError('无法解析日期')


def parse_int(value):
    if value is None:
        return 0
    s = str(value).strip()
    if s == '':
        return 0
    s = s.replace(',', '')
    if s.endswith('%'):
        s = s[:-1]
    try:
        return int(float(s))
    except Exception:
        return 0


def parse_float(value):
    if value is None:
        return 0.0
    s = str(value).strip()
    if s == '':
        return 0.0
    s = s.replace(',', '')
    is_percentage = s.endswith('%')
    if is_percentage:
        s = s[:-1]  # 去掉百分号
    try:
        float_val = float(s)
        if is_percentage:
            return float_val / 100.0  # 将百分比转换为小数
        return float_val
    except Exception:
        try:
            float_val = float(int(s))
            if is_percentage:
                return float_val / 100.0  # 将百分比转换为小数
            return float_val
        except Exception:
            return 0.0


def check_record_exists(record_id: str) -> bool:
    """检查记录是否已存在"""
    conn = get_connection()
    with conn:
        with conn.cursor() as cur:
            cur.execute('SELECT COUNT(*) as count FROM `gameData` WHERE `id` = %s', (record_id,))
            result = cur.fetchone()
            return result['count'] > 0


def import_rainbow_excel(file_path: str, game_id: str) -> Dict:
    """导入彩虹数据excel，返回详细的导入统计信息"""
    ensure_game_data_table()

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    header_row = [cell.value if cell.value is not None else '' for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    header_index: Dict[str, int] = {}
    for idx, title in enumerate(header_row):
        title_str = str(title).strip()
        if title_str in HEADER_TO_FIELD:
            header_index[HEADER_TO_FIELD[title_str]] = idx

    required = ['date', 'channel']
    missing = [k for k in required if k not in header_index]
    if missing:
        raise ValueError('缺少必要列: ' + ','.join(missing))

    # 准备批量插入的数据
    insert_data = []
    update_data = []
    
    # 先收集所有要处理的数据
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None:
            continue
        raw_date = row[header_index['date']]
        raw_channel = row[header_index['channel']]
        if raw_channel is None:
            continue
            
        date_val = parse_date(raw_date)
        channel_val = str(raw_channel).strip()
        row_id = generate_game_data_id(date_val, channel_val)

        # daNewUser 优先取"DA新增用户"，否则回退取"新增用户"
        if 'daNewUser' in header_index:
            da_new_user_val = parse_int(row[header_index.get('daNewUser')])
        elif 'newUser' in header_index:
            da_new_user_val = parse_int(row[header_index.get('newUser')])
        else:
            da_new_user_val = 0

        params = (
            row_id,
            game_id,
            date_val,
            '官包',
            channel_val,
            da_new_user_val,
            parse_int(row[header_index.get('newUser')]) if 'newUser' in header_index else 0,
            parse_int(row[header_index.get('active')]) if 'active' in header_index else 0,
            parse_int(row[header_index.get('payUser')]) if 'payUser' in header_index else 0,
            parse_float(row[header_index.get('gain')]) if 'gain' in header_index else 0.0,
            parse_float(row[header_index.get('ARPU')]) if 'ARPU' in header_index else 0.0,
            parse_float(row[header_index.get('ARPPU')]) if 'ARPPU' in header_index else 0.0,
            parse_float(row[header_index.get('DAY1')]) if 'DAY1' in header_index else 0.0,
            parse_float(row[header_index.get('DAY7')]) if 'DAY7' in header_index else 0.0,
        )
        
        # 检查记录是否已存在
        if check_record_exists(row_id):
            update_data.append(params)
        else:
            insert_data.append(params)

    # 批量执行插入和更新
    inserted_count = 0
    updated_count = 0
    
    conn = get_connection()
    with conn:
        with conn.cursor() as cur:
            # 批量插入新记录
            if insert_data:
                insert_sql = (
                    'INSERT INTO `gameData` '
                    '(`id`,`gameID`,`date`,`category`,`channel`,`daNewUser`,`newUser`,`active`,`payUser`,`gain`,`ARPU`,`ARPPU`,`DAY1`,`DAY7`) '
                    'VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)'
                )
                cur.executemany(insert_sql, insert_data)
                inserted_count = len(insert_data)
            
            # 批量更新已存在的记录
            if update_data:
                update_sql = (
                    'UPDATE `gameData` SET '
                    '`gameID`=%s,`date`=%s,`category`=%s,`channel`=%s,`daNewUser`=%s,`newUser`=%s,`active`=%s,'
                    '`payUser`=%s,`gain`=%s,`ARPU`=%s,`ARPPU`=%s,`DAY1`=%s,`DAY7`=%s '
                    'WHERE `id`=%s'
                )
                # 重新排列参数顺序：将id放在最后
                update_params = []
                for params in update_data:
                    # params[0]是id，需要移到最后
                    update_params.append(params[1:] + (params[0],))
                
                cur.executemany(update_sql, update_params)
                updated_count = len(update_data)
        
        conn.commit()

    return {
        'inserted': inserted_count,
        'updated': updated_count,
        'total': inserted_count + updated_count
    } 


def _parse_wg_date_str(date_str: str) -> datetime:
    """解析WG日期字符串，格式可能为 2025-09-02(二) -> 2025-09-02"""
    if isinstance(date_str, datetime):
        return date_str
    s = str(date_str).strip()
    if '(' in s:
        s = s.split('(', 1)[0].strip()
    # 复用已有解析器
    return parse_date(s)


def import_wg_excel(file_path: str, game_id: str, alias: str) -> Dict:
    """导入彩虹WG数据excel，按规则合并并写入数据库。

    规则：
    - 表头映射：
      日期->date，渠道名称->channel，"WG包" 固定为 category，
      新注册用户数->newUser，流水收入(元)->gain，付费用户数->payUser。
    - 如果【渠道名称】与下拉别名相同，则渠道=WeGame，否则= X8。
    - 将 X8 渠道按日期聚合：newUser/payUser/gain 求和，聚合后渠道= X8。
    - 日期如 2025-09-02(二) 需转为 2025-09-02。
    - 其余字段无则填 0 或 0.0。
    """
    ensure_game_data_table()

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    header_row = [cell.value if cell.value is not None else '' for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    header_index: Dict[str, int] = {}
    wg_header_map = {
        '日期': 'date',
        '渠道名称': 'channel',
        '新注册用户数': 'newUser',
        '流水收入(元)': 'gain',
        '付费用户数': 'payUser',
    }
    for idx, title in enumerate(header_row):
        title_str = str(title).strip()
        if title_str in wg_header_map:
            header_index[wg_header_map[title_str]] = idx

    required = ['date', 'channel']
    missing = [k for k in required if k not in header_index]
    if missing:
        raise ValueError('缺少必要列: ' + ','.join(missing))

    alias_norm = (alias or '').strip().lower()

    # 暂存行：WeGame 直接存，X8 先汇总
    x8_by_date = {}
    wegame_rows = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None:
            continue
        raw_date = row[header_index['date']]
        raw_channel = row[header_index['channel']]
        if raw_channel is None:
            continue

        date_val = _parse_wg_date_str(raw_date)
        channel_raw = str(raw_channel).strip()
        channel_norm = channel_raw.lower()
        mapped_channel = 'WeGame' if channel_norm == alias_norm and alias_norm != '' else 'X8'

        new_user = parse_int(row[header_index.get('newUser')]) if 'newUser' in header_index else 0
        pay_user = parse_int(row[header_index.get('payUser')]) if 'payUser' in header_index else 0
        gain_val = parse_float(row[header_index.get('gain')]) if 'gain' in header_index else 0.0

        if mapped_channel == 'X8':
            key = date_val.date()
            agg = x8_by_date.get(key)
            if not agg:
                x8_by_date[key] = {
                    'date': date_val,
                    'newUser': new_user,
                    'payUser': pay_user,
                    'gain': gain_val,
                }
            else:
                agg['newUser'] += new_user
                agg['payUser'] += pay_user
                agg['gain'] += gain_val
        else:
            wegame_rows.append({
                'date': date_val,
                'channel': 'WeGame',
                'newUser': new_user,
                'payUser': pay_user,
                'gain': gain_val,
            })

    # 组装最终插入/更新的数据
    insert_data = []
    update_data = []

    # WeGame 行
    for item in wegame_rows:
        row_id = generate_game_data_id(item['date'], item['channel'])
        params = (
            row_id,
            game_id,
            item['date'],
            'WG包',
            item['channel'],
            item['newUser'],  # daNewUser 使用 新增用户
            item['newUser'],
            0,  # active
            item['payUser'],
            item['gain'],
            0.0,  # ARPU
            0.0,  # ARPPU
            0.0,  # DAY1
            0.0,  # DAY7
        )
        if check_record_exists(row_id):
            update_data.append(params)
        else:
            insert_data.append(params)

    # X8 聚合行
    for key, agg in x8_by_date.items():
        date_val = agg['date']
        channel_val = 'X8'
        row_id = generate_game_data_id(date_val, channel_val)
        params = (
            row_id,
            game_id,
            date_val,
            'WG包',
            channel_val,
            agg['newUser'],
            agg['newUser'],
            0,
            agg['payUser'],
            agg['gain'],
            0.0,
            0.0,
            0.0,
            0.0,
        )
        if check_record_exists(row_id):
            update_data.append(params)
        else:
            insert_data.append(params)

    # 批量入库
    inserted_count = 0
    updated_count = 0
    conn = get_connection()
    with conn:
        with conn.cursor() as cur:
            if insert_data:
                insert_sql = (
                    'INSERT INTO `gameData` '
                    '(`id`,`gameID`,`date`,`category`,`channel`,`daNewUser`,`newUser`,`active`,`payUser`,`gain`,`ARPU`,`ARPPU`,`DAY1`,`DAY7`) '
                    'VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)'
                )
                cur.executemany(insert_sql, insert_data)
                inserted_count = len(insert_data)

            if update_data:
                update_sql = (
                    'UPDATE `gameData` SET '
                    '`gameID`=%s,`date`=%s,`category`=%s,`channel`=%s,`daNewUser`=%s,`newUser`=%s,`active`=%s,'
                    '`payUser`=%s,`gain`=%s,`ARPU`=%s,`ARPPU`=%s,`DAY1`=%s,`DAY7`=%s '
                    'WHERE `id`=%s'
                )
                update_params = []
                for params in update_data:
                    update_params.append(params[1:] + (params[0],))
                cur.executemany(update_sql, update_params)
                updated_count = len(update_data)
        conn.commit()

    return {
        'inserted': inserted_count,
        'updated': updated_count,
        'total': inserted_count + updated_count
    }


def import_wegame_excel(file_path: str, game_id: str) -> Dict:
    """导入WeGame数据excel（从第三行开始读），字段映射并批量入库/更新。

    表头映射：
    - 日期/Date -> date
    - 新增启动用户/New Launches by Users -> newUser
    - 启动用户数/Launches by Users -> active
    - 新进用户次日留存率/2nd Day Rentention for New Users -> DAY1

    固定：channel='WeGame', category='WG包'
    从第三行(min_row=3)开始读取。
    """
    ensure_game_data_table()

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    # 扫描前若干行寻找表头，并做鲁棒匹配（处理“日期\nDate”等情况）
    def normalize_title(x: str) -> str:
        s = (x or '').strip()
        s = s.replace('\n', ' ').replace('\r', ' ')
        s = ' '.join(s.split())  # collapse spaces
        return s.lower()

    header_index: Dict[str, int] = {}

    # 候选关键字集合（多语言、多写法）
    candidates = {
        'date': ['日期', 'date'],
        'newUser': ['新增启动用户', 'new launches by users', '新增用户'],
        'active': ['启动用户数', 'launches by users', '活跃用户'],
        'DAY1': ['新进用户次日留存率', '2nd day retention for new users', '2nd day rentention for new users', '次日留存'],
    }

    def try_map_header(titles):
        idx_map = {}
        used_indices = set()  # 记录已使用的列索引
        norm_titles = [normalize_title(str(t) if t is not None else '') for t in titles]
        
        # 第一轮：精确匹配
        for i, nt in enumerate(norm_titles):
            if not nt or i in used_indices:
                continue
            for field, keys in candidates.items():
                if field in idx_map:
                    continue
                for k in keys:
                    nk = normalize_title(k)
                    if nk and nk == nt:
                        idx_map[field] = i
                        used_indices.add(i)
                        break
                if field in idx_map:
                    break
        
        # 第二轮：包含匹配（仅对未匹配的字段）
        for i, nt in enumerate(norm_titles):
            if not nt or i in used_indices:
                continue
            for field, keys in candidates.items():
                if field in idx_map:
                    continue
                for k in keys:
                    nk = normalize_title(k)
                    if nk and (nk in nt or nt in nk):
                        idx_map[field] = i
                        used_indices.add(i)
                        break
                if field in idx_map:
                    break
        return idx_map

    best_map = {}
    best_row_idx = 1
    for r in range(1, 6):  # 扫描前5行
        row_vals = [cell.value if cell.value is not None else '' for cell in next(ws.iter_rows(min_row=r, max_row=r))]
        idx_map = try_map_header(row_vals)
        if len(idx_map) > len(best_map):
            best_map = idx_map
            best_row_idx = r
        if len(best_map) >= 3:  # 已找到大部分
            break
    header_index = best_map

    required = ['date']
    missing = [k for k in required if k not in header_index]
    if missing:
        raise ValueError('缺少必要列: ' + ','.join(missing))

    insert_data = []
    update_data = []

    # 数据从第三行开始读取
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row is None:
            continue
        raw_date = row[header_index['date']]
        if raw_date is None:
            continue
        # 日期可能带(周几)，做清洗
        date_val = _parse_wg_date_str(raw_date)
        channel_val = 'WeGame'
        category_val = 'WG包'

        new_user = parse_int(row[header_index.get('newUser')]) if 'newUser' in header_index else 0
        active_val = parse_int(row[header_index.get('active')]) if 'active' in header_index else 0
        day1_val = parse_float(row[header_index.get('DAY1')]) if 'DAY1' in header_index else 0.0

        row_id = generate_game_data_id(date_val, channel_val)

        if check_record_exists(row_id):
            # 仅更新Excel包含的字段；未包含的字段用旧值。且“新增用户”只影响 newUser，不改动 daNewUser。
            conn = get_connection()
            with conn:
                with conn.cursor() as cur:
                    cur.execute(
                        'SELECT `gameID`,`date`,`category`,`channel`,`daNewUser`,`newUser`,`active`,`payUser`,`gain`,`ARPU`,`ARPPU`,`DAY1`,`DAY7` '
                        'FROM `gameData` WHERE `id`=%s',
                        (row_id,)
                    )
                    old = cur.fetchone() or {}

            has_new_user = 'newUser' in header_index
            has_active = 'active' in header_index
            has_day1 = 'DAY1' in header_index

            merged_game_id = old.get('gameID', game_id)
            merged_date = old.get('date', date_val)  # 与row_id一致
            merged_category = old.get('category', category_val)
            merged_channel = old.get('channel', channel_val)
            merged_da_new_user = old.get('daNewUser', 0)  # 不随Excel的“新增用户”变化
            # 容错：如果Excel提供了新增用户但值为0，则沿用旧值，避免被0覆盖
            if has_new_user:
                merged_new_user = new_user if new_user != 0 else old.get('newUser', 0)
            else:
                merged_new_user = old.get('newUser', 0)
            # 如果Excel提供了活跃用户但值为0，则保留旧值
            if has_active:
                merged_active = active_val if active_val != 0 else old.get('active', 0)
            else:
                merged_active = old.get('active', 0)
            merged_pay_user = old.get('payUser', 0)
            merged_gain = old.get('gain', 0.0)
            merged_arpu = old.get('ARPU', 0.0)
            merged_arppu = old.get('ARPPU', 0.0)
            # 如果Excel提供了一留但值为0或0.0，则保留旧值
            if has_day1:
                merged_day1 = day1_val if day1_val != 0 and day1_val != 0.0 else old.get('DAY1', 0.0)
            else:
                merged_day1 = old.get('DAY1', 0.0)
            merged_day7 = old.get('DAY7', 0.0)

            params = (
                row_id,
                merged_game_id,
                merged_date,
                merged_category,
                merged_channel,
                merged_da_new_user,
                merged_new_user,
                merged_active,
                merged_pay_user,
                merged_gain,
                merged_arpu,
                merged_arppu,
                merged_day1,
                merged_day7,
            )
            update_data.append(params)
        else:
            # 插入：未提供的字段用默认值；daNewUser 使用 new_user
            params = (
                row_id,
                game_id,
                date_val,
                category_val,
                channel_val,
                new_user,                 # daNewUser 使用新增启动用户
                new_user,                 # newUser
                active_val,               # active
                0,                        # payUser (默认)
                0.0,                      # gain (默认)
                0.0,                      # ARPU
                0.0,                      # ARPPU
                day1_val,                 # DAY1
                0.0,                      # DAY7 (未提供)
            )
            insert_data.append(params)

    inserted_count = 0
    updated_count = 0
    conn = get_connection()
    with conn:
        with conn.cursor() as cur:
            if insert_data:
                insert_sql = (
                    'INSERT INTO `gameData` '
                    '(`id`,`gameID`,`date`,`category`,`channel`,`daNewUser`,`newUser`,`active`,`payUser`,`gain`,`ARPU`,`ARPPU`,`DAY1`,`DAY7`) '
                    'VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)'
                )
                cur.executemany(insert_sql, insert_data)
                inserted_count = len(insert_data)

            if update_data:
                update_sql = (
                    'UPDATE `gameData` SET '
                    '`gameID`=%s,`date`=%s,`category`=%s,`channel`=%s,`daNewUser`=%s,`newUser`=%s,`active`=%s,'
                    '`payUser`=%s,`gain`=%s,`ARPU`=%s,`ARPPU`=%s,`DAY1`=%s,`DAY7`=%s '
                    'WHERE `id`=%s'
                )
                update_params = []
                for params in update_data:
                    update_params.append(params[1:] + (params[0],))
                cur.executemany(update_sql, update_params)
                updated_count = len(update_data)
        conn.commit()

    return {
        'inserted': inserted_count,
        'updated': updated_count,
        'total': inserted_count + updated_count
    }